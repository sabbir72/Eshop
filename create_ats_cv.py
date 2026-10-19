#!/usr/bin/env python3
"""
Create ATS-optimized CV in Word format using python-docx
"""

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Installing python-docx...")
    import subprocess
    subprocess.check_call(['pip3', 'install', '--user', 'python-docx'])
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create document
doc = Document()

# Set margins (0.75 inch all sides for ATS compatibility)
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Helper function to add styled text
def add_heading(text, level=1):
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = heading.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(14) if level == 1 else Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return heading

def add_paragraph(text, bold=False, italic=False, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    return para

def add_bullet(text, size=11):
    para = doc.add_paragraph(text, style='List Bullet')
    run = para.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    return para

# NAME - Centered
name = doc.add_heading('TUHIN AHMED LIMA', level=0)
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name_run = name.runs[0]
name_run.font.name = 'Calibri'
name_run.font.size = Pt(18)
name_run.font.bold = True

# CONTACT INFO - Centered
contact1 = add_paragraph('Email: tuhin.lima@example.com | Phone: +880 1XXX-XXXXXX', size=10)
contact1.alignment = WD_ALIGN_PARAGRAPH.CENTER

contact2 = add_paragraph('LinkedIn: linkedin.com/in/tuhin-lima | GitHub: github.com/tuhin-lima', size=10)
contact2.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()  # Spacing

# PROFESSIONAL SUMMARY
add_heading('PROFESSIONAL SUMMARY', level=1)
add_paragraph('Results-driven Software Engineer with 5+ years of experience in full-stack development, specializing in web applications, mobile automation, and cloud technologies. Proven expertise in JavaScript, Python, React, Node.js, and Android automation frameworks. Strong background in building scalable applications, implementing CI/CD pipelines, and leading cross-functional teams.')
doc.add_paragraph()

# CORE COMPETENCIES
add_heading('CORE COMPETENCIES', level=1)
skills_para = doc.add_paragraph()
skills_para.add_run('Programming Languages: ').font.bold = True
skills_para.add_run('JavaScript, Python, Java, TypeScript, SQL, HTML5, CSS3\n')
skills_para.add_run('Frontend Technologies: ').font.bold = True
skills_para.add_run('React.js, Vue.js, Angular, Next.js, Redux, Tailwind CSS, Bootstrap\n')
skills_para.add_run('Backend Technologies: ').font.bold = True
skills_para.add_run('Node.js, Express.js, Django, Flask, REST APIs, GraphQL\n')
skills_para.add_run('Mobile Development: ').font.bold = True
skills_para.add_run('Android, React Native, Flutter, Appium, Mobile Automation\n')
skills_para.add_run('Databases: ').font.bold = True
skills_para.add_run('PostgreSQL, MySQL, MongoDB, Firebase, Redis, Supabase\n')
skills_para.add_run('Cloud & DevOps: ').font.bold = True
skills_para.add_run('AWS, Google Cloud, Docker, Kubernetes, CI/CD, Jenkins, GitHub Actions\n')
skills_para.add_run('Testing & Automation: ').font.bold = True
skills_para.add_run('Jest, Pytest, Selenium, Appium, UI Automation, Test-Driven Development')

for run in skills_para.runs:
    run.font.name = 'Calibri'
    run.font.size = Pt(11)

doc.add_paragraph()

# PROFESSIONAL EXPERIENCE
add_heading('PROFESSIONAL EXPERIENCE', level=1)

# Job 1
job1_title = add_paragraph('Senior Software Engineer', bold=True, size=12)
add_paragraph('ABC Tech Solutions, Dhaka | January 2023 - Present', italic=True, size=10)
add_bullet('Architected and developed enterprise-level e-commerce platform serving 50,000+ active users, resulting in 40% increase in sales conversion')
add_bullet('Led team of 5 developers in building scalable microservices using Node.js, Express, and PostgreSQL with 99.9% uptime')
add_bullet('Implemented CI/CD pipeline using GitHub Actions and Docker, reducing deployment time by 60%')
add_bullet('Integrated multiple payment gateways (bKash, SSLCommerz, Nagad) processing $500K+ monthly transactions')
add_bullet('Built React-based admin dashboard with real-time analytics, inventory management, and automated reporting')
add_bullet('Optimized database queries and implemented caching strategies, improving application response time by 45%')
doc.add_paragraph()

# Job 2
add_paragraph('Software Engineer', bold=True, size=12)
add_paragraph('XYZ Digital Agency, Dhaka | March 2021 - December 2022', italic=True, size=10)
add_bullet('Developed 15+ responsive web applications for clients across various industries using React and Vue.js')
add_bullet('Built RESTful APIs using Node.js and Express, serving mobile and web applications')
add_bullet('Implemented automated testing framework using Jest and Cypress, achieving 85% code coverage')
add_bullet('Created Android mobile automation scripts using Appium and Python for QA testing')
add_bullet('Integrated Firebase Authentication and Firestore for real-time data synchronization')
add_bullet('Optimized frontend performance achieving 90+ Google Lighthouse scores')
doc.add_paragraph()

# Job 3
add_paragraph('Junior Software Developer', bold=True, size=12)
add_paragraph('DEF Software Ltd., Dhaka | June 2019 - February 2021', italic=True, size=10)
add_bullet('Developed and maintained web applications using JavaScript, HTML5, CSS3, and jQuery')
add_bullet('Implemented responsive UI components using Bootstrap and custom CSS')
add_bullet('Performed bug fixes and feature enhancements based on client feedback')
add_bullet('Collaborated with senior developers on code reviews and pair programming sessions')
doc.add_paragraph()

# EDUCATION
add_heading('EDUCATION', level=1)
add_paragraph('Bachelor of Science in Computer Science and Engineering', bold=True, size=11)
add_paragraph('XYZ University, Dhaka, Bangladesh | Graduated: 2019 | CGPA: 3.75/4.00', size=10)
doc.add_paragraph()

# CERTIFICATIONS
add_heading('CERTIFICATIONS', level=1)
add_bullet('AWS Certified Solutions Architect - Associate | Amazon Web Services | 2024')
add_bullet('MongoDB Certified Developer | MongoDB University | 2023')
add_bullet('React - The Complete Guide | Udemy | 2022')
add_bullet('Mobile Automation with Appium | Test Automation University | 2023')
doc.add_paragraph()

# TECHNICAL PROJECTS
add_heading('TECHNICAL PROJECTS', level=1)
add_paragraph('Smart E-Commerce Management System', bold=True, size=11)
add_paragraph('Tech Stack: React, TypeScript, Node.js, Express, PostgreSQL, Prisma ORM, Firebase, Tailwind CSS', italic=True, size=10)
add_bullet('Built full-stack e-commerce platform with admin panel, inventory management, and order processing')
add_bullet('Implemented AI-powered shopping assistant using Google Gemini API for product recommendations')
add_bullet('Integrated multiple payment gateways with automatic reconciliation')
doc.add_paragraph()

add_paragraph('Android Mobile Automation Framework', bold=True, size=11)
add_paragraph('Tech Stack: Python, Appium, Selenium, Android SDK, ADB', italic=True, size=10)
add_bullet('Developed automated testing framework for Android applications using Appium and Python')
add_bullet('Created reusable test scripts for WhatsApp, browser, and e-commerce app automation')
add_bullet('Reduced manual testing time by 70% and increased test coverage by 50%')
doc.add_paragraph()

# LANGUAGES
add_heading('LANGUAGES', level=1)
lang_para = doc.add_paragraph()
lang_para.add_run('Bengali: ').font.bold = True
lang_para.add_run('Native  |  ')
lang_para.add_run('English: ').font.bold = True
lang_para.add_run('Fluent (Written and Spoken)')
for run in lang_para.runs:
    run.font.name = 'Calibri'
    run.font.size = Pt(11)

# Save document
output_path = '/home/eng-sabbir/Downloads/Tuhin_Ahmed_Lima_CV_ATS.docx'
doc.save(output_path)

print("✅ ATS-optimized CV created successfully!")
print(f"📄 File saved: {output_path}")
print("\n📋 ATS Optimization Features:")
print("  ✓ Simple layout without tables or columns")
print("  ✓ Standard fonts (Calibri)")
print("  ✓ Clear section headings")
print("  ✓ Bullet points for easy parsing")
print("  ✓ Keywords optimized for tech roles")
print("  ✓ Consistent formatting throughout")
